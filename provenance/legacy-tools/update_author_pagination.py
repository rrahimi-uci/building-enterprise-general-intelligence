from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.table import Table
import hashlib, json, re

DATA=Path('/data')
DOCX=DATA/'Building-Enterprise-Agentic-AI-Systems-Final.docx'
MD=DATA/'Building-Enterprise-Agentic-AI-Systems-Final.md'
ART=DATA/'v3-agent-cover-art.jpeg'
COVER=DATA/'Enterprise-Agentic-AI-Architecture-Guide-Cover-v3.1.png'

# Compose a print-ready cover using agent-focused artwork and visible author credit.
art=Image.open(ART).convert('RGB')
W,H=2550,3300
scale=max(W/art.width,H/art.height)
art=art.resize((int(art.width*scale),int(art.height*scale)),Image.Resampling.LANCZOS)
art=art.crop(((art.width-W)//2,(art.height-H)//2,(art.width+W)//2,(art.height+H)//2)).convert('RGBA')
overlay=Image.new('RGBA',(W,H),(0,0,0,0)); od=ImageDraw.Draw(overlay)
for y in range(H):
    top=max(0,1-y/1480); bottom=max(0,(y-2180)/1120)
    alpha=int(min(225,178*top+165*bottom))
    od.line((0,y,W,y),fill=(3,17,38,alpha))
art=Image.alpha_composite(art,overlay); d=ImageDraw.Draw(art)
reg='/usr/share/fonts/msttcorefonts/Arial.ttf'; bold='/usr/share/fonts/msttcorefonts/Arial_Bold.ttf'
if not Path(reg).exists():
    reg='/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf'; bold='/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf'
def F(n,b=False): return ImageFont.truetype(bold if b else reg,n)
# Kicker and title.
d.rounded_rectangle((165,145,1215,245),radius=22,fill=(24,104,210,225))
d.text((215,175),'ENTERPRISE ARCHITECTURE REFERENCE',font=F(45,True),fill='white')
y=360
for line in ['ENTERPRISE','AGENTIC AI','ARCHITECTURE']:
    d.text((165,y),line,font=F(150,True),fill='white',stroke_width=2,stroke_fill=(8,31,58)); y+=170
d.text((170,900),'& DEVELOPER PLATFORM GUIDE',font=F(63,True),fill=(81,220,234))
d.line((170,1010,1600,1010),fill=(224,154,50),width=10)
for i,line in enumerate(['Agent runtimes · Multi-agent systems · Skills & tools',
                         'Gateways · Security · Evaluation · Continuous learning',
                         'Governance · Deployment · Production operations']):
    d.text((175,1070+i*68),line,font=F(41),fill=(224,238,249))
# Author and edition lockup.
d.rounded_rectangle((155,2740,1325,2920),radius=26,fill=(5,25,52,225),outline=(81,220,234,255),width=4)
d.text((205,2780),'REZA RAHIMI, PhD',font=F(70,True),fill='white')
d.rounded_rectangle((155,2965,1090,3090),radius=20,fill=(5,25,52,225),outline=(224,154,50,255),width=4)
d.text((205,2998),'FINAL EDITION · v3.1',font=F(51,True),fill='white')
d.text((165,3160),'Enterprise architecture for governed, observable, and production-grade AI agents',font=F(34),fill=(205,225,241))
art.convert('RGB').save(COVER,quality=96)

# Update DOCX in place.
doc=Document(DOCX)
body=doc._body._element

# Replace the first embedded image (the cover) while preserving its layout and crop.
first_blip=next(body.iter(qn('a:blip')))
rid=first_blip.get(qn('r:embed'))
image_part=doc.part.related_parts[rid]
image_part._blob=COVER.read_bytes()

# Replace version labels without flattening paragraph formatting.
def replace_runs(paragraph):
    for run in paragraph.runs:
        if 'v3.0' in run.text:
            run.text=run.text.replace('v3.0','v3.1')
for p in doc.paragraphs: replace_runs(p)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs: replace_runs(p)
for section in doc.sections:
    for container in [section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer]:
        for p in container.paragraphs: replace_runs(p)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs: replace_runs(p)

# Add author to the accessible title page directly after the edition line.
if not any(p.text.strip()=='Reza Rahimi, PhD' for p in doc.paragraphs[:30]):
    edition=next(p for p in doc.paragraphs[:30] if 'edition' in p.text.lower() and 'v3.1' in p.text.lower())
    author=doc.add_paragraph()
    try: author.style='Subtitle'
    except KeyError: author.style='Normal'
    author.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=author.add_run('Reza Rahimi, PhD'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string('102A4C')
    edition._p.addnext(author._p)

# Page-break helpers.
def new_page_break_paragraph():
    p=OxmlElement('w:p'); r=OxmlElement('w:r'); br=OxmlElement('w:br'); br.set(qn('w:type'),'page'); r.append(br); p.append(r); return p

def is_page_break_paragraph(el):
    if el is None or el.tag!=qn('w:p'): return False
    return any(br.get(qn('w:type'))=='page' for br in el.iter(qn('w:br')))

def is_part_table(el):
    if el.tag!=qn('w:tbl'): return False
    try:
        t=Table(el,doc._body)
        return bool(t.rows and t.columns and t.cell(0,0).text.strip().startswith('PART '))
    except Exception:
        return False

# Make each Part divider a full standalone page: explicit page break before and after,
# tall navy panel, and vertically centered content.
for tbl_el in [e for e in list(body) if is_part_table(e)]:
    children=list(body); idx=children.index(tbl_el)
    if idx==0 or not is_page_break_paragraph(children[idx-1]):
        body.insert(idx,new_page_break_paragraph())
    children=list(body); idx=children.index(tbl_el)
    if idx+1>=len(children) or not is_page_break_paragraph(children[idx+1]):
        body.insert(idx+1,new_page_break_paragraph())
    t=Table(tbl_el,doc._body)
    t.autofit=False
    t.rows[0].height=Inches(8.45)
    t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
    t.cell(0,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

# The divider's explicit trailing page break starts its first chapter. Every
# later chapter/appendix receives pageBreakBefore so it starts on a fresh page.
chapter_count=0; waiting_after_part=False
for el in list(body):
    if is_part_table(el):
        waiting_after_part=True
        continue
    if el.tag!=qn('w:p'):
        continue
    p=Paragraph(el,doc._body)
    is_chapter=p.style.name=='Heading 1' and (bool(re.match(r'^\d{2} — ',p.text.strip())) or p.text.strip().startswith('Appendix '))
    if is_chapter:
        chapter_count+=1
        p.paragraph_format.page_break_before=not waiting_after_part
        p.paragraph_format.keep_with_next=True
        waiting_after_part=False

# Metadata.
doc.core_properties.author='Reza Rahimi, PhD'
doc.core_properties.last_modified_by='Reza Rahimi, PhD'
doc.core_properties.title='Enterprise Agentic AI Architecture & Developer Platform Guide — Final Edition v3.1'
doc.core_properties.comments='Final Edition v3.1: Reza Rahimi, PhD author credit; agent-focused cover; standalone Part pages; every chapter and appendix begins on a new page.'
doc.save(DOCX)

# Keep the canonical Markdown aligned for author/version metadata. Pagination and the
# rendered cover remain DOCX presentation concerns.
md=MD.read_text(encoding='utf-8')
md=md.replace('v3.0','v3.1')
if 'Reza Rahimi, PhD' not in md[:2000]:
    md=md.replace('Final Edition v3.1 · September 2026','Final Edition v3.1 · September 2026\n\n**Reza Rahimi, PhD**',1)
MD.write_text(md,encoding='utf-8')

manifest={
    'edition':'v3.1','author':'Reza Rahimi, PhD',
    'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    'markdown_sha256':hashlib.sha256(MD.read_bytes()).hexdigest(),
    'cover_sha256':hashlib.sha256(COVER.read_bytes()).hexdigest(),
    'chapter_and_appendix_headings':chapter_count,
    'part_dividers':len([e for e in list(body) if is_part_table(e)]),
    'docx_inline_images':len(doc.inline_shapes),
}
(DATA/'v3.1_preflight_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))
