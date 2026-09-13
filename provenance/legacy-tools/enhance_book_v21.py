from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import re
import shutil
import subprocess

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

DOCX = Path('/data/Building-Enterprise-Agentic-AI-Systems-Final.docx')
BACKUP = Path('/data/Building-Enterprise-Agentic-AI-Systems-Final-v2.0.docx')
MD = Path('/data/Building-Enterprise-Agentic-AI-Systems-Final.md')
ANTI = Path('/data/anti_patterns_v21.md').read_text(encoding='utf-8').strip()
CH23 = Path('/data/chapter23_v21.md').read_text(encoding='utf-8').strip()
EXTRACTED = Path('/data/Building-Enterprise-Agentic-AI-Systems-Final-extracted.md')
PART_DIR = Path('/data/final_parts')
PART_DIR.mkdir(parents=True, exist_ok=True)

if not BACKUP.exists():
    shutil.copy2(DOCX, BACKUP)
# Rebuild v2.1 from the immutable v2.0 release on every run.
shutil.copy2(BACKUP, DOCX)
doc = Document(DOCX)


def iter_paragraphs(document: Document):
    for p in document.paragraphs:
        yield p
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for section in document.sections:
        for container in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            for p in container.paragraphs:
                yield p
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p


def set_text(p: Paragraph, text: str) -> None:
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def add_inline_runs(p: Paragraph, text: str) -> None:
    pattern = re.compile(r'(\*\*.*?\*\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        else:
            run = p.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def insert_paragraph_before(anchor: Paragraph, text: str, style: str = 'Normal') -> Paragraph:
    p = doc.add_paragraph()
    try:
        p.style = style
    except KeyError:
        p.style = 'Normal'
    add_inline_runs(p, text)
    anchor._p.addprevious(p._p)
    return p


def insert_paragraph_after(anchor: Paragraph, text: str, style: str = 'Normal') -> Paragraph:
    p = doc.add_paragraph()
    try:
        p.style = style
    except KeyError:
        p.style = 'Normal'
    add_inline_runs(p, text)
    anchor._p.addnext(p._p)
    return p


def clean_cell_text(text: str) -> str:
    return re.sub(r'\*\*(.*?)\*\*', r'\1', text).replace('`', '').strip()


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 8.5) -> None:
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(clean_cell_text(text))
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def insert_table_before(anchor: Paragraph, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass
    table.autofit = True
    for ri, row in enumerate(rows):
        for ci in range(cols):
            value = row[ci] if ci < len(row) else ''
            set_cell_text(table.cell(ri, ci), value, bold=(ri == 0), size=8 if cols >= 5 else 8.5)
            if ri == 0:
                shade_cell(table.cell(ri, ci), 'D9EAF7')
    anchor._p.addprevious(table._tbl)


def insert_callout_before(anchor: Paragraph, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass
    set_cell_text(table.cell(0, 0), text, bold=False, size=9)
    shade_cell(table.cell(0, 0), 'EAF2F8')
    anchor._p.addprevious(table._tbl)


def insert_code_before(anchor: Paragraph, code: str) -> None:
    lines = code.splitlines() or ['']
    table = doc.add_table(rows=1, cols=2)
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass
    table.autofit = False
    table.columns[0].width = Inches(0.45)
    table.columns[1].width = Inches(6.65)
    left, right = table.rows[0].cells
    shade_cell(left, 'E9EEF3')
    shade_cell(right, 'F7F9FB')
    for cell in (left, right):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = OxmlElement('w:tcMar')
        for side, twips in [('top', '80'), ('left', '100'), ('bottom', '80'), ('right', '100')]:
            node = OxmlElement(f'w:{side}')
            node.set(qn('w:w'), twips)
            node.set(qn('w:type'), 'dxa')
            tc_mar.append(node)
        tc_pr.append(tc_mar)
    for idx, line in enumerate(lines):
        lp = left.paragraphs[0] if idx == 0 else left.add_paragraph()
        rp = right.paragraphs[0] if idx == 0 else right.add_paragraph()
        lp.clear(); rp.clear()
        for p in (lp, rp):
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
        lr = lp.add_run(str(idx + 1)); lr.font.name = 'Courier New'; lr.font.size = Pt(7); lr.font.color.rgb = RGBColor(90, 105, 120)
        rr = rp.add_run(line if line else ' '); rr.font.name = 'Courier New'; rr.font.size = Pt(7.5)
    anchor._p.addprevious(table._tbl)


def is_table_separator(line: str) -> bool:
    if not line.strip().startswith('|'):
        return False
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return bool(cells) and all(re.fullmatch(r':?-{3,}:?', c) for c in cells)


def split_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip('|').split('|')]


def parse_markdown(md: str):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('```'):
            language = line[3:].strip() or 'text'
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            yield ('code', language, '\n'.join(code_lines))
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            yield ('heading', len(m.group(1)), m.group(2).strip())
            i += 1
            continue
        if line.strip().startswith('|') and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header = split_table_row(line)
            i += 2
            data = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                data.append(split_table_row(lines[i]))
                i += 1
            if all(not c for c in header) and len(data) == 1 and len(data[0]) == 1:
                yield ('callout', clean_cell_text(data[0][0]))
            else:
                yield ('table', [header] + data)
            continue
        if line.startswith('> '):
            yield ('quote', line[2:].strip())
            i += 1
            continue
        if re.match(r'^[-*]\s+', line):
            yield ('bullet', re.sub(r'^[-*]\s+', '', line))
            i += 1
            continue
        if re.match(r'^\d+\.\s+', line):
            yield ('number', line)
            i += 1
            continue
        paragraph_lines = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip():
                break
            if nxt.startswith('#') or nxt.startswith('```') or nxt.startswith('> ') or nxt.strip().startswith('|') or re.match(r'^[-*]\s+', nxt) or re.match(r'^\d+\.\s+', nxt):
                break
            paragraph_lines.append(nxt.strip())
            i += 1
        yield ('paragraph', ' '.join(paragraph_lines))


def insert_markdown_before(anchor: Paragraph, md: str) -> None:
    for block in parse_markdown(md):
        kind = block[0]
        if kind == 'heading':
            p = insert_paragraph_before(anchor, block[2], f'Heading {block[1]}')
            if block[1] == 1:
                p.paragraph_format.page_break_before = True
        elif kind == 'paragraph':
            insert_paragraph_before(anchor, block[1], 'Normal')
        elif kind == 'quote':
            insert_paragraph_before(anchor, block[1], 'Quote')
        elif kind == 'bullet':
            insert_paragraph_before(anchor, block[1], 'List Bullet')
        elif kind == 'number':
            # Preserve source numbering explicitly; do not inherit document-wide numbering.
            insert_paragraph_before(anchor, block[1], 'Normal')
        elif kind == 'table':
            insert_table_before(anchor, block[1])
        elif kind == 'callout':
            insert_callout_before(anchor, block[1])
        elif kind == 'code':
            insert_code_before(anchor, block[2])


# Edition and front-matter updates.
for p in list(iter_paragraphs(doc)):
    text = p.text
    updated = text.replace('Final edition v2.0', 'Final edition v2.1').replace('FINAL EDITION · v2.0', 'FINAL EDITION · v2.1')
    updated = updated.replace('the 22 chapters and appendices', 'the 23 chapters and appendices')
    updated = updated.replace('1–3, 9–22', '1–3, 9–23').replace('9, 16, 18–22', '9, 16, 18–23')
    if updated != text:
        set_text(p, updated)

for p in doc.paragraphs:
    if p.text.startswith('All 17 review findings and all 10 assurance areas are integrated'):
        set_text(p, 'All 17 review findings and all 10 assurance areas are integrated into the 23 chapters and appendices. The v2.1 deep-review expansion adds a practical 16-pattern anti-pattern field guide and a complete continuous-learning, monitoring, and improvement architecture. Workload-specific conformance, privacy, security, legal, integration, recovery, and load tests remain deployment acceptance evidence—not missing book content.')
        break

# Add review improvements to the front summary table.
for table in doc.tables:
    if len(table.rows) >= 2 and len(table.columns) == 2 and table.cell(0, 0).text.strip() == 'Weakness' and table.cell(0, 1).text.strip() == 'Improvement':
        existing = {r.cells[0].text.strip() for r in table.rows}
        additions = [
            ('Anti-pattern guidance too terse', 'Sixteen practical field-guide cards covering causes, manifestations, detection, impact, remediation, design comparisons, and metrics.'),
            ('Post-deployment learning fragmented', 'A complete Observe → Measure → Diagnose → Recommend → Validate → Deploy → Monitor architecture with automation and approval boundaries.'),
        ]
        for left, right in additions:
            if left not in existing:
                row = table.add_row()
                set_cell_text(row.cells[0], left, False, 8.5)
                set_cell_text(row.cells[1], right, False, 8.5)
        break

# Add chapter 23 to every explicit Part V contents list.
chapter_22_label = '22 Cost, Latency, Capacity, and Reliability Optimization'
for p in [p for p in doc.paragraphs if p.text.strip() == chapter_22_label]:
    nxt = p._p.getnext()
    nxt_text = ''.join(nxt.itertext()).strip() if nxt is not None else ''
    if '23 Continuous Learning' not in nxt_text:
        insert_paragraph_after(p, '23 Continuous Learning, Monitoring, and Improvement', p.style.name)

# Replace terse Appendix A anti-pattern index with the field guide.
anti_start = next(p for p in doc.paragraphs if p.text.strip() == 'Anti-pattern index')
appendix_b = next(p for p in doc.paragraphs if p.text.strip().startswith('Appendix B —'))
body = doc._body._element
children = list(body)
start_idx = children.index(anti_start._p)
end_idx = children.index(appendix_b._p)
for element in children[start_idx:end_idx]:
    body.remove(element)
insert_markdown_before(appendix_b, ANTI)

# Add Chapter 23 before Appendix A. Remove inherited manual page-break
# paragraphs so Heading 1 creates exactly one new page, not a blank page.
appendix_a = next(p for p in doc.paragraphs if p.text.strip().startswith('Appendix A —'))
prev = appendix_a._p.getprevious()
while prev is not None:
    prev_text = ''.join(prev.itertext()).strip()
    has_break = bool(prev.findall('.//' + qn('w:br')))
    if prev.tag != qn('w:p') or prev_text or not has_break:
        break
    prior = prev.getprevious()
    doc._body._element.remove(prev)
    prev = prior
appendix_a.paragraph_format.page_break_before = True
insert_markdown_before(appendix_a, CH23)

# Record the v2.1 expansion in the closure log.
appendix_m = next(p for p in doc.paragraphs if p.text.strip().startswith('Appendix M —'))
insert_paragraph_before(appendix_m, 'v2.1 deep-review expansion', 'Heading 2')
insert_paragraph_before(appendix_m, '✓ Appendix A now contains a 16-pattern practical anti-pattern field guide with detection, impact, remediation, design comparisons, and operational signals.', 'Normal')
insert_paragraph_before(appendix_m, '✓ Chapter 23 now defines the closed continuous-improvement architecture, evaluation framework landscape, metrics, layer-specific optimization paths, automation boundaries, and regression-suite lifecycle.', 'Normal')

# Add primary-source ledger entries before Appendix N.
appendix_n = next(p for p in doc.paragraphs if p.text.strip().startswith('Appendix N —'))
for entry in [
    'DEEPEVAL-EVAL · FRAMEWORK. Agent trajectory, component, and RAG evaluation metrics. https://deepeval.com/docs/metrics-introduction',
    'MLFLOW-CONTINUOUS · FRAMEWORK. Evaluation-driven development, datasets, feedback, judges, and production monitoring. https://mlflow.org/docs/latest/genai/eval-monitor/',
    'MLFLOW-FEEDBACK · FRAMEWORK. Trace-linked automated, programmatic, and human feedback with attribution and override history. https://mlflow.org/docs/latest/genai/assessments/feedback/',
    'OTEL-GENAI-REPO · STANDARD / DEVELOPMENT. Separate OpenTelemetry GenAI conventions repository for spans, metrics, events, MCP, and provider mappings. https://github.com/open-telemetry/semantic-conventions-genai',
    'PHOENIX-OPENINFERENCE · FRAMEWORK. OpenTelemetry/OpenInference tracing, evaluation, datasets, experiments, prompt versioning, and replay. https://arize.com/docs/phoenix',
    'RAGAS-AGENT · FRAMEWORK. Agent/tool and RAG evaluation metrics, including tool-call accuracy/F1 and agent-goal accuracy. https://docs.ragas.io/en/latest/concepts/metrics/available_metrics/agents',
    'LANGSMITH-EVAL · FRAMEWORK. Offline and online evaluation with production traces, datasets, and regression workflows. https://docs.langchain.com/langsmith/evaluation',
]:
    insert_paragraph_before(appendix_n, entry, 'List Bullet')

# Update metadata and save.
doc.core_properties.title = 'Building Enterprise Agentic AI Systems — Final Edition v2.1'
doc.core_properties.subject = 'Grounded final architecture, anti-pattern, and continuous-improvement reference'
doc.core_properties.comments = 'Final edition v2.1; expanded anti-pattern field guide and continuous learning, monitoring, and improvement chapter; DOCX and Markdown aligned.'
doc.core_properties.modified = datetime.now(timezone.utc)
doc.save(DOCX)

# Convert the updated DOCX to canonical Markdown.
with EXTRACTED.open('w', encoding='utf-8') as f:
    proc = subprocess.run(['python3', '-m', 'markitdown', str(DOCX)], stdout=f, stderr=subprocess.PIPE, text=True)
if proc.returncode != 0:
    raise SystemExit(proc.stderr)
md = EXTRACTED.read_text(encoding='utf-8')

# Convert every syntax-highlighted two-column DOCX code table into a fenced block.
doc_for_extract = Document(DOCX)
code_tables = []
for table in doc_for_extract.tables:
    if len(table.columns) != 2 or len(table.rows) != 1:
        continue
    number_lines = [p.text.strip() for p in table.rows[0].cells[0].paragraphs if p.text.strip()]
    if not number_lines or number_lines[0] != '1':
        continue
    if number_lines != [str(i + 1) for i in range(len(number_lines))]:
        continue
    code_lines = [p.text.rstrip() for p in table.rows[0].cells[1].paragraphs]
    code_tables.append('\n'.join('' if line.isspace() else line for line in code_lines).rstrip())

lines = md.splitlines()
out_lines = []
code_index = 0
language_map = {'PYTHON': 'python', 'BASH': 'bash', 'JSON': 'json', 'SQL': 'sql', 'TEXT': 'text'}
i = 0
while i < len(lines):
    if i + 2 < len(lines) and lines[i].strip() == '|  |  |' and lines[i + 1].strip() == '| --- | --- |' and out_lines:
        prev = next((x.strip() for x in reversed(out_lines) if x.strip()), '')
        if prev.endswith('· implementation excerpt'):
            if code_index >= len(code_tables):
                raise RuntimeError('more code-table markers than extracted code tables')
            label = prev.split('·', 1)[0].strip().upper()
            out_lines.extend([f'```{language_map.get(label, "text")}', code_tables[code_index], '```'])
            code_index += 1
            i += 3
            continue
    out_lines.append(lines[i])
    i += 1
if code_index != len(code_tables):
    raise RuntimeError(f'converted {code_index} of {len(code_tables)} code tables')
md = '\n'.join(out_lines) + '\n'

# Replace embedded binary images with stable Notion markers while preserving captions/tables.
image_counter = 0
def repl_img(match):
    global image_counter
    image_counter += 1
    if image_counter == 1:
        return '> Cover artwork is embedded in the DOCX edition.'
    return f'> Architecture figure {image_counter - 1} is embedded in the DOCX; the immediately following caption and blueprint table are canonical.'
md = re.sub(r'!\[\]\(data:image/[^\n]+\)', repl_img, md)
MD.write_text(md, encoding='utf-8')

# Split the canonical source for internal parity checks.
markers = [
    ('01-front-matter.md', None, '# 01 — From an LLM Call to a Production System'),
    ('02-part-i.md', '# 01 — From an LLM Call to a Production System', '# 04 — ReAct: Bounded Reason–Act–Observe Loops'),
    ('03-part-ii.md', '# 04 — ReAct: Bounded Reason–Act–Observe Loops', '# 10 — Skills as Versioned Capabilities'),
    ('04-part-iii.md', '# 10 — Skills as Versioned Capabilities', '# 14 — Context Engineering, Memory, and State'),
    ('05-part-iv.md', '# 14 — Context Engineering, Memory, and State', '# 18 — Observability and Tracing with MLflow'),
    ('06-part-v.md', '# 18 — Observability and Tracing with MLflow', '# Appendix A — Pattern selection and anti-patterns'),
    ('07-appendices.md', '# Appendix A — Pattern selection and anti-patterns', None),
]
for filename, start, end in markers:
    s = 0 if start is None else md.index(start)
    e = len(md) if end is None else md.index(end)
    (PART_DIR / filename).write_text(md[s:e].strip() + '\n', encoding='utf-8')

def normalize(text: str) -> str:
    text = re.sub(r'> (Cover artwork|Architecture figure).*?canonical\.', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

parts = '\n'.join(p.read_text(encoding='utf-8') for p in sorted(PART_DIR.glob('*.md')))
manifest = {
    'edition': 'v2.1',
    'as_of': '2026-09-12',
    'docx_sha256': hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    'canonical_text_sha256': hashlib.sha256(normalize(md).encode()).hexdigest(),
    'rejoined_text_sha256': hashlib.sha256(normalize(parts).encode()).hexdigest(),
    'normalized_text_match': normalize(md) == normalize(parts),
    'chapters': len(re.findall(r'^# (?:0[1-9]|1[0-9]|2[0-3]) — ', md, flags=re.M)),
    'appendices': len(re.findall(r'^# Appendix [A-O] — ', md, flags=re.M)),
    'anti_patterns': len(re.findall(r'^### AP-\d{2} — ', md, flags=re.M)),
    'code_blocks': len(re.findall(r'^```(?:python|bash|json|sql|text)$', md, flags=re.M)),
    'figure_assets_in_docx': len(doc_for_extract.inline_shapes),
    'figure_markers_in_markdown': len(re.findall(r'^> (?:Cover artwork|Architecture figure)', md, flags=re.M)),
    'continuous_loop_stages_present': all(stage in md for stage in ['Observe', 'Measure', 'Diagnose', 'Recommend', 'Validate', 'Deploy', 'Monitor']),
    'frameworks_present': all(name in md for name in ['DeepEval', 'MLflow', 'OpenTelemetry', 'Phoenix', 'Ragas', 'LangSmith']),
    'requested_layers_present': all(name in md.lower() for name in ['prompts', 'tools', 'skills', 'memory', 'context management', 'routing', 'agent workflow', 'model selection', 'deterministic logic', 'overall behavior']),
}
Path('/data/final_alignment_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
print(json.dumps(manifest, indent=2))
