from __future__ import annotations
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import hashlib, json, math, re, shutil, subprocess, textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.table import Table

BASE=Path('/data/Building-Enterprise-Agentic-AI-Systems-Final.docx')
BACKUP=Path('/data/Building-Enterprise-Agentic-AI-Systems-Final-v2.1.docx')
OUT_DOCX=BASE
OUT_MD=Path('/data/Building-Enterprise-Agentic-AI-Systems-Final.md')
SRC_DIR=Path('/data/v3_sources')
DIAGRAM_DIR=Path('/data/v3_diagrams')
COVER_ART=Path('/data/v3-cover-art.jpeg')
COVER_FINAL=Path('/data/v3-cover-final.png')
EXTRACTED=Path('/data/v3-extracted.md')
PART_DIR=Path('/data/final_parts_v3'); PART_DIR.mkdir(exist_ok=True)

if not BACKUP.exists(): shutil.copy2(BASE,BACKUP)
doc=Document(BACKUP)
plan=json.loads((SRC_DIR/'plan.json').read_text())
parts=plan['parts']; chapters=plan['chapters']; part_chapters=plan['part_chapters']

NAVY='102A4C'; NAVY2='173E68'; BLUE='1769D2'; CYAN='20B8D8'; TEAL='118A84'; AMBER='E09A32'; RED='C94A4A'; LIGHT='EEF4FA'; WHITE='FFFFFF'; INK='12233F'

# ---------- cover ----------
def make_cover():
    art=Image.open(COVER_ART).convert('RGB')
    W,H=2550,3300
    scale=max(W/art.width,H/art.height)
    art=art.resize((int(art.width*scale),int(art.height*scale)),Image.Resampling.LANCZOS)
    x=(art.width-W)//2; y=(art.height-H)//2
    im=art.crop((x,y,x+W,y+H)).convert('RGBA')
    overlay=Image.new('RGBA',(W,H),(0,0,0,0)); od=ImageDraw.Draw(overlay)
    for yy in range(H):
        top=max(0,1-yy/1500); bottom=max(0,(yy-2200)/1100)
        a=int(175*top+150*bottom)
        od.line((0,yy,W,yy),fill=(4,18,40,min(220,a)))
    im=Image.alpha_composite(im,overlay)
    d=ImageDraw.Draw(im)
    reg='/usr/share/fonts/msttcorefonts/Arial.ttf'; bold='/usr/share/fonts/msttcorefonts/Arial_Bold.ttf'
    if not Path(reg).exists(): reg='/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf'; bold='/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf'
    def f(s,b=False): return ImageFont.truetype(bold if b else reg,s)
    d.rounded_rectangle((170,155,1130,245),radius=20,fill=(25,105,210,220))
    d.text((220,177),'ENTERPRISE ARCHITECTURE REFERENCE',font=f(44,True),fill='white')
    title=['ENTERPRISE','AGENTIC AI','ARCHITECTURE']
    yy=370
    for line in title:
        d.text((170,yy),line,font=f(150,True),fill='white',stroke_width=1,stroke_fill=(12,38,72)); yy+=170
    d.text((178,910),'& DEVELOPER PLATFORM GUIDE',font=f(62,True),fill=(70,220,235))
    d.line((175,1020,1500,1020),fill=(224,154,50),width=10)
    sub=['Reference architecture · SDK & developer experience',
         'Patterns · Security · Gateways · Continuous learning',
         'Governance · Deployment · Production operations']
    yy=1080
    for line in sub:
        d.text((180,yy),line,font=f(42),fill=(225,238,250)); yy+=66
    d.rounded_rectangle((170,2900,1040,3020),radius=20,fill=(8,29,58,220),outline=(58,196,214,255),width=4)
    d.text((215,2934),'FINAL EDITION · v3.0',font=f(50,True),fill='white')
    d.text((175,3120),'Architecture for executives, architects, platform teams, and senior engineers',font=f(34),fill=(205,224,242))
    im.convert('RGB').save(COVER_FINAL,quality=95)
make_cover()

# ---------- capture reusable v2.1 chapter/appendix XML ----------
body=doc._body._element
children=list(body)
heading_positions=[]
for i,e in enumerate(children):
    if e.tag!=qn('w:p'): continue
    p=Paragraph(e,doc._body)
    if p.style.name=='Heading 1' and (re.match(r'^\d{2} — ',p.text) or p.text.startswith('Appendix ')):
        heading_positions.append((i,p.text))

blocks={}
for pos,(start,title) in enumerate(heading_positions):
    end=heading_positions[pos+1][0] if pos+1<len(heading_positions) else len(children)-1
    block=[deepcopy(x) for x in children[start:end]]
    # Old part dividers sit at the end of the preceding chapter. Remove them.
    cut=None
    for j,e in enumerate(block):
        if e.tag==qn('w:tbl'):
            try:
                t=Table(e,doc._body)
                if t.cell(0,0).text.strip().startswith('PART '): cut=j; break
            except Exception: pass
    if cut is not None:
        while cut>0 and block[cut-1].tag==qn('w:p') and not ''.join(block[cut-1].itertext()).strip(): cut-=1
        block=block[:cut]
    if title.startswith('Appendix '): key='APP-'+title.split()[1]
    else: key=title[:2]
    blocks[key]=block

# Clear body while retaining final section properties.
for e in list(body):
    if e.tag!=qn('w:sectPr'): body.remove(e)

# ---------- styles and layout ----------
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.28)
styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(9.6); styles['Normal'].font.color.rgb=RGBColor.from_string(INK)
styles['Normal'].paragraph_format.space_after=Pt(5); styles['Normal'].paragraph_format.line_spacing=1.06
for name,size,color,space in [('Title',30,NAVY,14),('Subtitle',14,TEAL,8),('Heading 1',24,NAVY,14),('Heading 2',16,BLUE,9),('Heading 3',12,NAVY2,7),('Heading 4',10,TEAL,5)]:
    try:
        st=styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=name!='Subtitle'; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(space); st.paragraph_format.space_after=Pt(space/2)
    except KeyError: pass
try:
    styles['Caption'].font.name='Arial'; styles['Caption'].font.size=Pt(8.5); styles['Caption'].font.italic=True; styles['Caption'].font.color.rgb=RGBColor.from_string('4A6078')
except KeyError: pass

# Header/footer
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.clear(); hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=hp.add_run('ENTERPRISE AGENTIC AI ARCHITECTURE & DEVELOPER PLATFORM GUIDE'); r.font.name='Arial'; r.font.size=Pt(7.5); r.font.bold=True; r.font.color.rgb=RGBColor.from_string('4A6078')
    fp=section.footer.paragraphs[0]; fp.clear(); fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=fp.add_run('FINAL EDITION · v3.0    |    PAGE '); rr.font.name='Arial'; rr.font.size=Pt(7.5); rr.font.color.rgb=RGBColor.from_string('4A6078')
    fld_begin=OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    fld_sep=OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'),'separate')
    txt=OxmlElement('w:t'); txt.text='1'
    fld_end=OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'),'end')
    run=OxmlElement('w:r'); run.extend([fld_begin,instr,fld_sep,txt,fld_end]); fp._p.append(run)
    section.different_first_page_header_footer=True
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()

# ---------- helpers ----------
def shade_cell(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell,top=70,left=90,bottom=70,right=90):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',left),('bottom',bottom),('end',right)]:
        node=tcMar.find(qn(f'w:{m}'))
        if node is None: node=OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def add_inline(p,text):
    pat=re.compile(r'(\*\*.*?\*\*|`[^`]+`)'); pos=0
    for m in pat.finditer(text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        tok=m.group(0)
        if tok.startswith('**'):
            r=p.add_run(tok[2:-2]); r.bold=True
        else:
            r=p.add_run(tok[1:-1]); r.font.name='Courier New'; r.font.size=Pt(8.8); r.font.color.rgb=RGBColor.from_string('405A75')
        pos=m.end()
    if pos<len(text): p.add_run(text[pos:])

def clean_md(text):
    return re.sub(r'\*\*(.*?)\*\*',r'\1',text).replace('`','').strip()

def add_para(text='',style='Normal',align=None):
    p=doc.add_paragraph(style=style); add_inline(p,text)
    if align is not None: p.alignment=align
    return p

def add_page_break():
    p=doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); return p

def add_table(rows,header=True,small=False):
    cols=max(len(r) for r in rows); t=doc.add_table(rows=len(rows),cols=cols)
    try: t.style='Table Grid'
    except: pass
    t.autofit=True
    for ri,row in enumerate(rows):
        for ci in range(cols):
            cell=t.cell(ri,ci); cell.text=''; set_cell_margins(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(clean_md(row[ci] if ci<len(row) else '')); r.font.name='Arial'; r.font.size=Pt(7.7 if small or cols>=5 else 8.2)
            if ri==0 and header:
                r.bold=True; r.font.color.rgb=RGBColor(255,255,255); shade_cell(cell,NAVY)
            elif ri%2==0: shade_cell(cell,'F3F7FB')
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

def add_code(code):
    lines=code.splitlines() or ['']; t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; t.autofit=False
    t.columns[0].width=Inches(.42); t.columns[1].width=Inches(6.45)
    left,right=t.rows[0].cells; shade_cell(left,'E5ECF3'); shade_cell(right,'F7F9FC')
    for c in (left,right): set_cell_margins(c,60,85,60,85)
    for i,line in enumerate(lines):
        lp=left.paragraphs[0] if i==0 else left.add_paragraph(); rp=right.paragraphs[0] if i==0 else right.add_paragraph(); lp.clear(); rp.clear()
        for p in (lp,rp): p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
        r=lp.add_run(str(i+1)); r.font.name='Courier New'; r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string('6A7E92')
        r=rp.add_run(line if line else ' '); r.font.name='Courier New'; r.font.size=Pt(7.3); r.font.color.rgb=RGBColor.from_string('19324D')
    return t

def add_diagram(key,caption):
    path=DIAGRAM_DIR/f'{key}.png'
    if not path.exists(): raise FileNotFoundError(path)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path),width=Inches(6.85))
    c=doc.add_paragraph(style='Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.add_run(caption)

def add_static_toc():
    part_lookup={p[0]:(p[1],p[2]) for p in parts}; ch_lookup={int(c[0]):c[1] for c in chapters}
    for roman in [p[0] for p in parts]:
        title,desc=part_lookup[roman]
        add_para(f'Part {roman} — {title}','Heading 2')
        add_para(desc,'Normal')
        for n in part_chapters[roman]: add_para(f'{n:02d}  {ch_lookup[n]}','Normal')
    add_para('Appendices A–O — Decision aids, runbooks, contracts, frameworks, assurance, and release acceptance','Heading 2')

def sep_row(line):
    cells=[c.strip() for c in line.strip().strip('|').split('|')]
    return bool(cells) and all(re.fullmatch(r':?-{3,}:?',c) for c in cells)

def split_row(line): return [c.strip() for c in line.strip().strip('|').split('|')]

def parse_md(md):
    lines=md.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        dm=re.fullmatch(r'\[\[DIAGRAM:([^|\]]+)\|(.+)\]\]',line.strip())
        if dm: yield ('diagram',dm.group(1),dm.group(2)); i+=1; continue
        if line.strip()=='[[TOC]]': yield ('toc',); i+=1; continue
        if line.startswith('```'):
            lang=line[3:].strip() or 'text'; code=[]; i+=1
            while i<len(lines) and not lines[i].startswith('```'): code.append(lines[i]); i+=1
            i+=1; yield ('code',lang,'\n'.join(code)); continue
        m=re.match(r'^(#{1,4})\s+(.*)$',line)
        if m: yield ('heading',len(m.group(1)),m.group(2).strip()); i+=1; continue
        if line.strip().startswith('|') and i+1<len(lines) and sep_row(lines[i+1]):
            rows=[split_row(line)]; i+=2
            while i<len(lines) and lines[i].strip().startswith('|'): rows.append(split_row(lines[i])); i+=1
            yield ('table',rows); continue
        if re.match(r'^[-*]\s+\[[ xX]\]\s+',line):
            checked='☑' if re.match(r'^[-*]\s+\[[xX]\]',line) else '☐'; text=re.sub(r'^[-*]\s+\[[ xX]\]\s+','',line); yield ('bullet',f'{checked}  {text}'); i+=1; continue
        if re.match(r'^[-*]\s+',line): yield ('bullet',re.sub(r'^[-*]\s+','',line)); i+=1; continue
        if re.match(r'^\d+\.\s+',line): yield ('number',line); i+=1; continue
        if line.startswith('> '): yield ('quote',line[2:].strip()); i+=1; continue
        pp=[line.strip()]; i+=1
        while i<len(lines):
            n=lines[i]
            if not n.strip(): break
            if n.startswith(('#','```','> ')) or n.strip().startswith('|') or re.match(r'^[-*]\s+',n) or re.match(r'^\d+\.\s+',n) or n.strip().startswith('[[DIAGRAM:') or n.strip()=='[[TOC]]': break
            pp.append(n.strip()); i+=1
        yield ('para',' '.join(pp))

def add_markdown(md,h1_page_break=True):
    first_h1=True
    for b in parse_md(md):
        kind=b[0]
        if kind=='heading':
            p=add_para(b[2],f'Heading {b[1]}')
            if b[1]==1: p.paragraph_format.page_break_before=h1_page_break if first_h1 else True; first_h1=False
        elif kind=='para': add_para(b[1])
        elif kind=='bullet': add_para(b[1],'List Bullet')
        elif kind=='number': add_para(b[1],'Normal')
        elif kind=='quote':
            t=add_table([['',b[1]]],header=False); shade_cell(t.cell(0,0),AMBER); shade_cell(t.cell(0,1),'FFF7E8')
        elif kind=='table': add_table(b[1],header=True,small=len(b[1][0])>=5)
        elif kind=='code': add_code(b[2])
        elif kind=='diagram': add_diagram(b[1],b[2])
        elif kind=='toc': add_static_toc()

def add_part_divider(roman,title,desc):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; t.autofit=False; cell=t.cell(0,0); shade_cell(cell,NAVY); set_cell_margins(cell,260,260,260,260)
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f'PART {roman}'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string('62D7E7')
    p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(title.upper()); r.bold=True; r.font.name='Arial'; r.font.size=Pt(28); r.font.color.rgb=RGBColor(255,255,255)
    p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(desc); r.font.name='Arial'; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string('C7DDF2')
    add_page_break()

def append_clone(old_key,new_title,page_break):
    els=[deepcopy(e) for e in blocks[old_key]]
    first=None
    for e in els:
        body.insert(len(body)-1,e)
        if first is None and e.tag==qn('w:p'): first=Paragraph(e,doc._body)
    if first is None: raise RuntimeError(old_key)
    first.clear(); first.style='Heading 1'; first.add_run(new_title); first.paragraph_format.page_break_before=page_break
    return first

# ---------- assemble ----------
# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.add_run().add_picture(str(COVER_FINAL),width=Inches(7.05))
add_page_break()
# Text title page supports accessible Markdown and navigation.
p=add_para('Enterprise General Intelligence: Architecting Governed Agentic Systems for the Enterprise','Title',WD_ALIGN_PARAGRAPH.CENTER); p.paragraph_format.space_before=Pt(110)
p=add_para('Reference architecture, platform SDK, patterns, security, gateways, continuous learning, governance, and production operations','Subtitle',WD_ALIGN_PARAGRAPH.CENTER)
add_para('FINAL EDITION · v3.0 · September 2026','Heading 2',WD_ALIGN_PARAGRAPH.CENTER)
add_para('A structured enterprise reference for executives, architects, platform teams, security leaders, SREs, and senior engineers.','Normal',WD_ALIGN_PARAGRAPH.CENTER)
add_page_break()
add_markdown((SRC_DIR/'00-front.md').read_text(),h1_page_break=False)

part_map={p[0]:(p[1],p[2]) for p in parts}; chapter_map={int(c[0]):c for c in chapters}
for roman in [p[0] for p in parts]:
    title,desc=part_map[roman]; add_part_divider(roman,title,desc)
    for j,n in enumerate(part_chapters[roman]):
        num,title,kind,source=chapter_map[n]
        if kind=='old': append_clone(source,f'{num:02d} — {title}',page_break=(j>0))
        else: add_markdown((SRC_DIR/source).read_text(),h1_page_break=(j>0))

# Appendix divider and appendices
add_part_divider('REFERENCE','Appendices','Decision aids, runbooks, contracts, frameworks, assurance evidence, and release acceptance.')
add_markdown((SRC_DIR/'appendix-a.md').read_text(),h1_page_break=False)
for letter in 'BCDEFGHIJKLMNO':
    block=blocks[f'APP-{letter}']; old_title=Paragraph(block[0],doc._body).text
    append_clone(f'APP-{letter}',old_title,page_break=True)

# ---------- global reference/version updates ----------
chapter_number_map={1:1,2:2,3:3,4:10,5:11,6:12,7:13,8:14,9:15,10:8,11:9,12:18,13:19,14:22,15:23,16:24,17:25,18:26,19:27,20:5,21:30,22:31,23:28}
ref_re=re.compile(r'\b(Chapter|chapter|Chapters|chapters)\s+(\d{1,2})(?:([–-])(\d{1,2}))?')
def remap_refs(text):
    def repl(m):
        a=chapter_number_map.get(int(m.group(2)),int(m.group(2)))
        if m.group(4):
            b=chapter_number_map.get(int(m.group(4)),int(m.group(4)))
            return f'{m.group(1)} {a}{m.group(3)}{b}'
        return f'{m.group(1)} {a}'
    text=ref_re.sub(repl,text)
    return (text.replace('Final edition v2.1','Final edition v3.0')
        .replace('FINAL EDITION · v2.1','FINAL EDITION · v3.0')
        .replace('v2.1 deep-review expansion','v3.0 architecture and developer-platform restructuring')
        .replace('Appendix A now contains a 16-pattern practical anti-pattern field guide','Chapter 17 now contains a 16-pattern practical anti-pattern catalog')
        .replace('the closed continuous-improvement architecture','the ten-stage continuous-learning and improvement architecture')
        .replace('the 23 chapters and appendices','the 31 chapters and appendices')
        .replace('22 chapters','31 chapters').replace('23 chapters','31 chapters'))

def replace_p_text(p,new):
    if new==p.text: return
    p.clear(); p.add_run(new)
for p in doc.paragraphs: replace_p_text(p,remap_refs(p.text))
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: replace_p_text(p,remap_refs(p.text))

# Add framework-abstraction note to Chapter 3.
for p in doc.paragraphs:
    if p.text=='03 — Agent Runtime and Framework Abstraction':
        note=doc.add_paragraph(); note.style='Normal'; add_inline(note,'This chapter uses the OpenAI Agents SDK as a concrete adapter example. Enterprise applications should compile platform-owned contracts into this or another framework through conformance-tested adapters; Chapters 4–7 define the portable architecture.')
        p._p.addnext(note._p); break

# Complete the v3 restructuring closure log.
for p in doc.paragraphs:
    if p.text.startswith('Appendix M —'):
        for text in [
            '✓ Added the enterprise reference architecture, platform boundaries, Agentic Developer Platform and SDK, developer experience, LLM Gateway, security/guardrails, and governance chapters.',
            '✓ Reorganized the guide into twelve parts with a hierarchical table of contents and a Foundations-to-Production narrative.',
            '✓ Standardized pattern and anti-pattern cards as Problem → Anti-Pattern → Why It Fails → Recommended Pattern → Architecture → Enterprise Example → Detection Metrics → Trade-offs.',
            '✓ Added sixteen threat cards with attack scenario, enterprise impact, detection, prevention, runtime controls, and recovery, plus defense-in-depth and trust-boundary diagrams.',
            '✓ Added a new editorial cover and eighteen architecture, sequence, lifecycle, decision, security, and operating-model diagrams; each visual teaches a named responsibility or control flow.'
        ]:
            np=doc.add_paragraph(); np.style='Normal'; add_inline(np,text); p._p.addprevious(np._p)
        break

# Improve all conventional tables, preserve code and one-cell callouts.
for t in doc.tables:
    if len(t.columns)==1: continue
    is_code=False
    if len(t.columns)==2 and len(t.rows)==1:
        nums=[p.text.strip() for p in t.cell(0,0).paragraphs if p.text.strip()]
        is_code=bool(nums and nums[0]=='1' and nums==[str(i+1) for i in range(len(nums))])
    if is_code: continue
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            set_cell_margins(c)
            if ri==0:
                shade_cell(c,NAVY)
                for p in c.paragraphs:
                    for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.name='Arial'; r.font.size=Pt(8)
            elif ri%2==0: shade_cell(c,'F3F7FB')
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs:
                    if r.font.size is None or r.font.size.pt>9: r.font.size=Pt(8)
                    r.font.name='Arial'

# Metadata
doc.core_properties.title='Enterprise General Intelligence: Architecting Governed Agentic Systems for the Enterprise — Final Edition v3.0'
doc.core_properties.subject='Enterprise reference architecture, developer platform SDK, patterns, security, gateways, continuous learning, governance, and production operations'
doc.core_properties.comments='Final edition v3.0; 31-chapter restructured architecture and developer platform guide; generated with aligned canonical Markdown.'
doc.core_properties.modified=datetime.now(timezone.utc)
doc.save(OUT_DOCX)

# ---------- canonical Markdown extraction ----------
with EXTRACTED.open('w',encoding='utf-8') as f:
    r=subprocess.run(['python3','-m','markitdown',str(OUT_DOCX)],stdout=f,stderr=subprocess.PIPE,text=True)
if r.returncode: raise SystemExit(r.stderr)
md=EXTRACTED.read_text(encoding='utf-8')

# Convert syntax-highlighted two-column code tables to fenced code.
d2=Document(OUT_DOCX); code_tables=[]
for t in d2.tables:
    if len(t.columns)!=2 or len(t.rows)!=1: continue
    nums=[p.text.strip() for p in t.cell(0,0).paragraphs if p.text.strip()]
    if not nums or nums[0]!='1' or nums!=[str(i+1) for i in range(len(nums))]: continue
    lines=[p.text.rstrip() for p in t.cell(0,1).paragraphs]
    code_tables.append('\n'.join('' if x.isspace() else x for x in lines).rstrip())
lines=md.splitlines(); out=[]; ci=0; langmap={'PYTHON':'python','BASH':'bash','JSON':'json','SQL':'sql','TEXT':'text'}; i=0
while i<len(lines):
    if i+2<len(lines) and lines[i].strip()=='|  |  |' and lines[i+1].strip()=='| --- | --- |' and out:
        prev=next((x.strip() for x in reversed(out) if x.strip()),'')
        if prev.endswith('· implementation excerpt'):
            if ci>=len(code_tables): raise RuntimeError('code table marker mismatch')
            label=prev.split('·',1)[0].strip().upper(); out.extend([f'```{langmap.get(label,"text")}',code_tables[ci],'```']); ci+=1; i+=3; continue
    out.append(lines[i]); i+=1
if ci!=len(code_tables): raise RuntimeError(f'converted {ci}/{len(code_tables)} code tables')
md='\n'.join(out)+'\n'

# Replace embedded images with stable semantic markers. Captions remain immediately after.
image_index=0
def img_repl(m):
    global image_index
    image_index+=1
    if image_index==1: return '> Cover artwork is embedded in the DOCX edition.'
    return f'> Teaching figure {image_index-1} is embedded in the DOCX; its caption and surrounding architecture content are canonical.'
md=re.sub(r'!\[\]\(data:image/[^\n]+\)',img_repl,md)
OUT_MD.write_text(md,encoding='utf-8')

# Split by parts for parity and review.
for p in PART_DIR.glob('*.md'): p.unlink()
part_starts=[]
for roman,title,desc in parts:
    idx=md.find(f'**PART {roman}**')
    if idx<0: idx=md.find(f'PART {roman}')
    if idx>=0:
        table_start=md.rfind('\n|  |\n| --- |\n',0,idx)
        idx=(table_start+1) if table_start>=0 else md.rfind('\n',0,idx)+1
    part_starts.append((roman,title,idx))
app_idx=md.find('**PART REFERENCE**')
if app_idx<0: app_idx=md.find('PART REFERENCE')
app_table_start=md.rfind('\n|  |\n| --- |\n',0,app_idx)
app_start=(app_table_start+1) if app_table_start>=0 else md.rfind('\n',0,app_idx)+1
for k,(roman,title,start) in enumerate(part_starts):
    if start<0: raise RuntimeError(f'missing part {roman}')
    end=part_starts[k+1][2] if k+1<len(part_starts) else app_start
    (PART_DIR/f'{k+1:02d}-part-{roman.lower()}.md').write_text(md[start:end].strip()+'\n',encoding='utf-8')
front_end=part_starts[0][2]
(PART_DIR/'00-front.md').write_text(md[:front_end].strip()+'\n',encoding='utf-8')
(PART_DIR/'13-appendices.md').write_text(md[app_start:].strip()+'\n',encoding='utf-8')

# QA manifest
def norm(x):
    x=re.sub(r'> (?:Cover artwork|Teaching figure).*?canonical\.', '', x)
    return re.sub(r'\s+',' ',x).strip()
rejoined='\n'.join(p.read_text(encoding='utf-8') for p in sorted(PART_DIR.glob('*.md')))
required_threats=['direct prompt injection','indirect prompt injection','jailbreaks','malicious context','retrieval/rag poisoning','memory poisoning','tool misuse','excessive agency','privilege escalation','data exfiltration','sensitive-data leakage','insecure agent-to-agent communication','malicious mcp servers','unsafe code execution','identity spoofing','supply-chain attacks']
required_gateway=['provider abstraction','model routing','authentication and authorization','policy enforcement','pii and secrets protection','rate limiting','token','cost','caching','retries','fallbacks','content filtering','model/version governance','logging and tracing','evaluation hooks']
required_platform=['agent runtime','workflow / orchestration','tool sdk','mcp integration','agent-to-agent communication','context / rag','memory','skills','identity & permissions','guardrails','llm gateway','evaluation','observability / tracing','durable execution','human-in-the-loop','deployment & versioning']
low=md.lower()
manifest={
 'edition':'v3.0','as_of':'2026-09-12',
 'docx_sha256':hashlib.sha256(OUT_DOCX.read_bytes()).hexdigest(),
 'markdown_sha256':hashlib.sha256(OUT_MD.read_bytes()).hexdigest(),
 'canonical_normalized_sha256':hashlib.sha256(norm(md).encode()).hexdigest(),
 'rejoined_normalized_sha256':hashlib.sha256(norm(rejoined).encode()).hexdigest(),
 'normalized_parts_match':norm(md)==norm(rejoined),
 'chapters':len(re.findall(r'^# (?:0[1-9]|[12][0-9]|3[01]) — ',md,re.M)),
 'appendices':len(re.findall(r'^# Appendix [A-O] — ',md,re.M)),
 'parts':len(re.findall(r'PART (?:I|II|III|IV|V|VI|VII|VIII|IX|X|XI|XII)\b',md)),
 'pattern_cards':len(re.findall(r'^### P-\d{2} — ',md,re.M)),
 'anti_pattern_cards':len(re.findall(r'^### AP-\d{2} — ',md,re.M)),
 'security_threats':len(re.findall(r'^### T-\d{2} — ',md,re.M)),
 'code_blocks':len(re.findall(r'^```(?:python|bash|json|sql|text)$',md,re.M)),
 'docx_inline_images':len(d2.inline_shapes),
 'markdown_image_markers':len(re.findall(r'^> (?:Cover artwork|Teaching figure)',md,re.M)),
 'learning_stages':{x:x.lower() in low for x in ['Observe','Measure','Detect','Diagnose','Recommend','Validate','Approve','Deploy','Monitor','Learn']},
 'threat_coverage':{x:x in low for x in required_threats},
 'gateway_coverage':{x:x in low for x in required_gateway},
 'platform_coverage':{x:x in low for x in required_platform},
 'stale_versions':{'v2.1':md.count('v2.1'),'v2.0':md.count('v2.0')},
}
Path('/data/final_alignment_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))
