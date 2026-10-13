#!/usr/bin/env python3
"""Deterministically build the publication DOCX from canonical Markdown.

Supported canonical constructs:
  headings H1-H4, paragraphs, bullet/number lists, blockquotes, GFM pipe
  tables, fenced code, inline bold/italic/code/links, [Rxx] citations,
  [[COVER]], and [[DIAGRAM:id|caption]].

The script owns all DOCX layout and style decisions. Do not hand-format release
artifacts; change book/manuscript.md, config, assets, or this renderer.
"""
from __future__ import annotations
from pathlib import Path
import book_index
import callouts
from datetime import datetime,timezone
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import argparse,hashlib,json,re,subprocess,sys,tempfile
from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches,Pt,RGBColor

FIXED_DT=datetime(2026,9,13,12,0,0,tzinfo=timezone.utc)
# The index tracker for the build in progress (tools/book_index.py decides where
# marks go; this renderer only turns them into Word XE fields).
INDEX=None

def load(root:Path):
    return json.loads((root/'config/book.json').read_text(encoding='utf-8'))

def color(v): return RGBColor.from_string(v)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def cell_margins(cell,top=70,left=90,bottom=70,right=90):
    tcPr=cell._tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for name,val in [('top',top),('start',left),('bottom',bottom),('end',right)]:
        n=mar.find(qn('w:'+name))
        if n is None: n=OxmlElement('w:'+name); mar.append(n)
        n.set(qn('w:w'),str(val)); n.set(qn('w:type'),'dxa')

def repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)

def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))

def add_hyperlink(p,text,url,hex_color='1769D2',underline=True):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),hex_color); rPr.append(c)
    if underline:
        u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u)
    r.append(rPr); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def configure(doc:Document,cfg:dict):
    C=cfg['design']['colors']; font=cfg['design']['fonts']; m=cfg['design']['margins_inches']
    sec=doc.sections[0]; sec.top_margin=Inches(m['top']); sec.bottom_margin=Inches(m['bottom']); sec.left_margin=Inches(m['left']); sec.right_margin=Inches(m['right']); sec.header_distance=Inches(.25); sec.footer_distance=Inches(.28); sec.different_first_page_header_footer=True
    styles=doc.styles
    normal=styles['Normal']; normal.font.name=font['sans']; normal.font.size=Pt(10); normal.font.color.rgb=color(C['ink']); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1666666667
    # Word stores complex-script size separately even when the book is English;
    # keep the named Normal style identical to the reviewed publication.
    normal_rpr=normal._element.get_or_add_rPr(); szcs=normal_rpr.find(qn('w:szCs'))
    if szcs is None: szcs=OxmlElement('w:szCs'); normal_rpr.append(szcs)
    szcs.set(qn('w:val'),'20')
    specs={'Title':(30,C['navy'],True,14,7),'Subtitle':(14,C['teal'],False,8,5),'Heading 1':(24,C['navy'],True,14,7),'Heading 2':(16,C['blue'],True,9,5),'Heading 3':(12,C['navy2'],True,7,4),'Heading 4':(10,C['teal'],True,5,3)}
    for name,(size,clr,bold,before,after) in specs.items():
        st=styles[name]; st.font.name=font['sans']; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=color(clr); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    cap=styles['Caption']; cap.font.name=font['sans']; cap.font.size=Pt(8.5); cap.font.italic=True; cap.font.color.rgb=color('4A6078'); cap.paragraph_format.space_before=Pt(2); cap.paragraph_format.space_after=Pt(6)
    for name,base in [('Book Callout','Normal'),('Book Evidence','Normal'),('Book Reference','Normal'),('Book Code Label','Normal'),('Book Part','Title')]:
        if name not in [s.name for s in styles]: styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=styles[name]; st.base_style=styles[base]; st.font.name=font['sans'] if name!='Book Code Label' else font['mono']
    styles['Book Evidence'].font.color.rgb=color(C['navy2']); styles['Book Evidence'].font.italic=True
    styles['Book Reference'].font.size=Pt(8.6); styles['Book Reference'].paragraph_format.left_indent=Inches(.18); styles['Book Reference'].paragraph_format.first_line_indent=Inches(-.18)
    styles['Book Code Label'].font.size=Pt(7.5); styles['Book Code Label'].font.bold=True; styles['Book Code Label'].font.color.rgb=color(C['white'])
    styles['Book Part'].font.size=Pt(28); styles['Book Part'].font.bold=True; styles['Book Part'].font.color.rgb=color(C['white'])
    # Header and footer.
    hp=sec.header.paragraphs[0]; hp.clear(); hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=hp.add_run(cfg['book']['title'].upper()); r.font.name=font['sans']; r.font.size=Pt(7.2); r.font.bold=True; r.font.color.rgb=color('4A6078')
    fp=sec.footer.paragraphs[0]; fp.clear(); fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run(f"BOOK v{cfg['book']['version']}    |    PAGE "); r.font.name=font['sans']; r.font.size=Pt(7.5); r.font.color.rgb=color('4A6078')
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '; sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); txt=OxmlElement('w:t'); txt.text='1'; end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); rr=OxmlElement('w:r'); rr.extend([begin,instr,sep,txt,end]); fp._p.append(rr)
    sec.first_page_header.paragraphs[0].clear(); sec.first_page_footer.paragraphs[0].clear()

def split_row(line):
    text=line.strip().strip('|'); cells=[]; buf=[]; esc=False; code=False
    for ch in text:
        if esc: buf.append(ch); esc=False; continue
        if ch=='\\': esc=True; buf.append(ch); continue
        if ch=='`': code=not code; buf.append(ch); continue
        if ch=='|' and not code: cells.append(''.join(buf).strip().replace('\\|','|')); buf=[]
        else: buf.append(ch)
    cells.append(''.join(buf).strip().replace('\\|','|')); return cells

def sep_row(line):
    cs=split_row(line); return bool(cs) and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cs)

def list_item(lines,i,first):
    # A hard-wrapped item continues on indented lines; join them so the item
    # renders as one list paragraph instead of a bullet plus stray paragraphs.
    parts=[first.strip()]; i+=1
    while i<len(lines) and lines[i][:1] in (' ','\t') and lines[i].strip() and not re.match(r'^\s*([-*]|\d+\.)\s+',lines[i]):
        parts.append(lines[i].strip()); i+=1
    return ' '.join(parts),i

def part_intro(raw):
    """A Part page's text: lead and closing paragraphs plus an optional list."""
    blocks=[]; para=[]
    def flush():
        if para: blocks.append(('para',' '.join(para))); para.clear()
    for line in raw:
        if not line.strip(): flush(); continue
        if re.match(r'^[-*]\s+',line): flush(); blocks.append(('bullet',re.sub(r'^[-*]\s+','',line).strip())); continue
        if line[:1] in (' ','\t') and blocks and blocks[-1][0]=='bullet' and not para:
            blocks[-1]=('bullet',blocks[-1][1]+' '+line.strip()); continue
        para.append(line.strip())
    flush(); return blocks

def parse(md):
    lines=md.splitlines(); i=0
    # Optional YAML front matter.
    if lines and lines[0].strip()=='---':
        i=1
        while i<len(lines) and lines[i].strip()!='---': i+=1
        i+=1
    while i<len(lines):
        line=lines[i]; stripped=line.strip()
        if not stripped: i+=1; continue
        if stripped=='[[COVER]]': yield ('cover',); i+=1; continue
        dm=re.fullmatch(r'\[\[DIAGRAM:([^|\]]+)\|(.+)\]\]',stripped)
        if dm: yield ('diagram',dm.group(1),dm.group(2)); i+=1; continue
        pm=re.match(r'^# PART (.+?) — (.+)$',line)
        if pm:
            ident,title=pm.groups(); i+=1
            while i<len(lines) and not lines[i].strip(): i+=1
            raw=[]
            while i<len(lines) and not lines[i].startswith('# '):
                raw.append(lines[i]); i+=1
            yield ('part',ident,title,part_intro(raw)); continue
        if line.startswith('```'):
            lang=line[3:].strip() or 'text'; code=[]; i+=1
            while i<len(lines) and not lines[i].startswith('```'): code.append(lines[i]); i+=1
            if i>=len(lines): raise ValueError('Unclosed code fence')
            i+=1; yield ('code',lang,'\n'.join(code)); continue
        hm=re.match(r'^(#{1,4})\s+(.+)$',line)
        if hm: yield ('heading',len(hm.group(1)),hm.group(2).strip()); i+=1; continue
        if stripped.startswith('|') and i+1<len(lines) and sep_row(lines[i+1]):
            rows=[split_row(line)]; i+=2
            while i<len(lines) and lines[i].strip().startswith('|'): rows.append(split_row(lines[i])); i+=1
            yield ('table',rows); continue
        if re.match(r'^[-*]\s+\[[ xX]\]\s+',line):
            checked='☑' if re.match(r'^[-*]\s+\[[xX]\]',line) else '☐'; yield ('bullet',checked+' '+re.sub(r'^[-*]\s+\[[ xX]\]\s+','',line)); i+=1; continue
        if re.match(r'^[-*]\s+',line): text,i=list_item(lines,i,re.sub(r'^[-*]\s+','',line)); yield ('bullet',text); continue
        if re.match(r'^\d+\.\s+',line): text,i=list_item(lines,i,re.sub(r'^\d+\.\s+','',line)); yield ('number',text); continue
        if line.startswith('> '): yield ('quote',line[2:].strip()); i+=1; continue
        para=[stripped]; i+=1
        while i<len(lines):
            n=lines[i]; ns=n.strip()
            if not ns: break
            if n.startswith(('#','```','> ')) or ns.startswith('|') or ns.startswith('[[DIAGRAM:') or ns=='[[COVER]]' or re.match(r'^[-*]\s+',n) or re.match(r'^\d+\.\s+',n): break
            para.append(ns); i+=1
        yield ('para',' '.join(para))

# Markdown backslash escapes (artifact\_id, a \< b) are source syntax, not text.
# Strip them everywhere except inside code spans, matching build_latex.py.
def md_unescape(text):
    return re.sub(r'\\([\\`*_{}\[\]()#+.!|<>\-])', r'\1', text)

def add_inline(p,text,cfg,bold_override=False):
    C=cfg['design']['colors']; mono=cfg['design']['fonts']['mono']
    token=re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|\[R\d{2}\]|(?<!\*)\*[^*\n]+\*)')
    pos=0
    for m in token.finditer(text):
        if m.start()>pos:
            r=p.add_run(md_unescape(text[pos:m.start()])); r.bold=bold_override or None
        t=m.group(0)
        if t.startswith('**'):
            r=p.add_run(md_unescape(t[2:-2])); r.bold=True
        elif t.startswith('`'):
            r=p.add_run(t[1:-1]); r.font.name=mono; r.font.size=Pt(8.7); r.font.color.rgb=color('405A75'); r.font.highlight_color=None
        elif re.match(r'^\[[^\]]+\]\(https?://',t):
            mm=re.match(r'^\[([^\]]+)\]\((https?://[^)]+)\)$',t); add_hyperlink(p,mm.group(1),mm.group(2),C['blue'])
        elif re.fullmatch(r'\[R\d{2}\]',t):
            r=p.add_run(t); r.bold=True; r.font.color.rgb=color(C['blue'])
        elif t.startswith('*'):
            r=p.add_run(md_unescape(t[1:-1])); r.italic=True
        pos=m.end()
    if pos<len(text):
        r=p.add_run(md_unescape(text[pos:])); r.bold=bold_override or None

def field_run(p,kind=None,instr=None):
    r=OxmlElement('w:r')
    if kind:
        f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),kind); r.append(f)
    else:
        t=OxmlElement('w:instrText'); t.set(qn('xml:space'),'preserve'); t.text=instr; r.append(t)
    return r

def mark_index(p,text):
    """Append a Word XE field for each term book_index marks in this text."""
    if INDEX is None: return
    for term,defining in INDEX.marks(text):
        instr=' XE "'+term.replace('"','')+'"'+(' \\b' if defining else '')+' '
        p._p.extend([field_run(p,'begin'),field_run(p,instr=instr),field_run(p,'end')])

def open_index_field(p):
    """Start Word's INDEX field before this paragraph's text. Its cached result
    is the Markdown index (chapter locators) until Word updates the field."""
    runs=[field_run(p,'begin'),field_run(p,instr=' INDEX \\h "A" \\e ", " '),field_run(p,'separate')]
    anchor=p._p.find(qn('w:pPr'))
    for i,r in enumerate(runs): p._p.insert((list(p._p).index(anchor)+1 if anchor is not None else 0)+i,r)

# Every numbered list restarts at 1. The List Number style carries one shared
# numbering instance, so each new list gets its own instance with a start override.
def new_numbered_list(doc):
    numbering=doc.part.numbering_part.numbering_definitions._numbering
    style_num=int(doc.styles['List Number'].element.pPr.numPr.numId.val)
    abstract=numbering.num_having_numId(style_num).abstractNumId.val
    num=numbering.add_num(abstract); num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId

def set_num(p,num_id):
    numPr=p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val=0; numPr.get_or_add_numId().val=num_id

def add_paragraph(doc,text,cfg,style='Normal',align=None):
    p=doc.add_paragraph(style=style); add_inline(p,text,cfg)
    if align is not None: p.alignment=align
    return p

def add_part(doc,ident,title,intro,cfg):
    C=cfg['design']['colors']
    # A one-point anchor paragraph starts the page: unlike a paragraph holding a break, it
    # cannot spill onto an empty page when the previous page is full.
    anchor=doc.add_paragraph(); anchor.paragraph_format.page_break_before=True; anchor.paragraph_format.space_before=Pt(0); anchor.paragraph_format.space_after=Pt(0)
    anchor.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY; anchor.paragraph_format.line_spacing=Pt(1)
    t=doc.add_table(rows=1,cols=1); t.autofit=False; row=t.rows[0]; row.height=Inches(8.45); row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY; cell=t.cell(0,0); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(cell,C['navy']); cell_margins(cell,300,320,300,320)
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('PART '+ident); r.font.name=cfg['design']['fonts']['sans']; r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=color(C['cyan'])
    p=cell.add_paragraph(style='Book Part'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title.upper())
    # The intro is set ragged-right in a centred measure, so a lead paragraph and
    # a chapter list read as one block; centred multi-line text is hard to read.
    for n,(kind,text) in enumerate(intro):
        p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Inches(.75); p.paragraph_format.right_indent=Inches(.75)
        p.paragraph_format.space_before=Pt(18 if n==0 else (3 if kind=='bullet' else 9)); p.paragraph_format.space_after=Pt(0)
        if kind=='bullet':
            p.paragraph_format.left_indent=Inches(1.0); p.paragraph_format.first_line_indent=Inches(-.2)
            r=p.add_run('•\t'); r.font.color.rgb=color(C['cyan']); r.font.size=Pt(11)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(1.0))
        add_inline(p,text,cfg)
        for r in p.runs:
            r.font.name=cfg['design']['fonts']['sans']; r.font.size=Pt(11)
            if r.text!='•\t': r.font.color.rgb=color(C['white'] if r.bold else 'D8E8F5')
    # No trailing break paragraph: whatever follows starts its own page with page_break_before,
    # so a page that is already full cannot push an empty page into the book.

def add_code_run(p,text,lang,cfg):
    # Lightweight deterministic lexical colors; source text is never changed.
    C=cfg['design']['colors']; kw=set('class def return if else elif for while try except finally with as import from async await raise yield True False None and or not in is lambda match case switch const let function public private new'.split())
    pieces=re.split(r'("[^"\n]*"|\'[^\'\n]*\'|\b\w+\b)',text)
    for piece in pieces:
        r=p.add_run(piece if piece else ''); r.font.name=cfg['design']['fonts']['mono']; r.font.size=Pt(7.25)
        if piece in kw: r.font.bold=True; r.font.color.rgb=color(C['blue'])
        elif (piece.startswith('"') and piece.endswith('"')) or (piece.startswith("'") and piece.endswith("'")): r.font.color.rgb=color(C['teal'])
        elif piece.strip().startswith(('#','//')): r.font.color.rgb=color('6A7E92')
        else: r.font.color.rgb=color('19324D')

def add_code(doc,lang,code,cfg,spacer=True):
    C=cfg['design']['colors']; t=doc.add_table(rows=1,cols=2); t.autofit=False; t.columns[0].width=Inches(.42); t.columns[1].width=Inches(6.42)
    h=t.rows[0].cells[0]; h.merge(t.rows[0].cells[1]); shade(h,C['navy']); cell_margins(h,45,90,45,90); p=h.paragraphs[0]; p.style='Book Code Label'; p.add_run('🐍' if (lang or 'text').lower() == 'python' else (lang or 'text').upper())
    for i,line in enumerate(code.splitlines() or [''],1):
        cells=t.add_row().cells; prevent_row_split(t.rows[-1]); shade(cells[0],'E5ECF3'); shade(cells[1],'F7F9FC'); cell_margins(cells[0],15,55,15,55); cell_margins(cells[1],15,80,15,80)
        lp,rp=cells[0].paragraphs[0],cells[1].paragraphs[0]
        for p in (lp,rp): p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
        rr=lp.add_run(str(i)); rr.font.name=cfg['design']['fonts']['mono']; rr.font.size=Pt(7); rr.font.color.rgb=color('6A7E92'); add_code_run(rp,line or ' ',lang,cfg)
    if spacer: doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_table(doc,rows,cfg,spacer=True):
    C=cfg['design']['colors']; cols=max(len(r) for r in rows)
    # One-column two-row Markdown tables are semantic callouts.
    if cols==1 and rows and not rows[0][0].strip() and len(rows)>=2:
        text=' '.join(r[0] for r in rows[1:] if r and r[0].strip())
        # The reviewed DOCX retains the separator after the bold callout kind
        # as a second whitespace character. It is layout-neutral in Word and
        # does not alter the canonical Markdown or XeLaTeX text.
        text=re.sub(r'(\*\*[^*]+\*\*) (?=\S)',r'\1  ',text)
        # The shared taxonomy picks the shading, so the DOCX and XeLaTeX editions
        # cannot drift apart. Every fill here is a pale ground: 'Book Callout'
        # inherits Normal's dark ink, and shading these cells navy -- as this did
        # in earlier drafts -- put #12233F text on #102A4C at a ratio of 1.09:1.
        _env, fill, _title = callouts.classify(text)
        t=doc.add_table(rows=1,cols=1); cell=t.cell(0,0); shade(cell,fill); cell_margins(cell,120,160,120,160)
        p=cell.paragraphs[0]; p.style='Book Callout'
        # The marker stays inline here: a DOCX callout has no title bar, so the
        # bold marker is what names the box. LaTeX strips it because the tcolorbox
        # title already carries it.
        add_inline(p,text,cfg); mark_index(p,text)
        if spacer: doc.add_paragraph().paragraph_format.space_after=Pt(1)
        return
    t=doc.add_table(rows=len(rows),cols=cols); t.style='Table Grid'; t.autofit=True
    size=6.8 if cols>=7 else (7.3 if cols>=5 else 8.1)
    for ri,row in enumerate(rows):
        tr=t.rows[ri]; prevent_row_split(tr)
        if ri==0: repeat_table_header(tr)
        for ci in range(cols):
            cell=t.cell(ri,ci); cell.text=''; cell_margins(cell,65,80,65,80); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
            val=row[ci] if ci<len(row) else ''; add_inline(p,val,cfg,bold_override=ri==0)
            if ri: mark_index(p,val)
            for run in p.runs: run.font.name=cfg['design']['fonts']['sans']; run.font.size=Pt(size)
            if ri==0:
                shade(cell,C['navy']);
                for run in p.runs: run.font.color.rgb=color(C['white']); run.bold=True
            elif ri%2==0: shade(cell,'F3F7FB')
    if spacer: doc.add_paragraph().paragraph_format.space_after=Pt(1)

def chapter_cover_path(root,number):
    spec=root/'config/chapter-covers.json'
    if not spec.exists(): return None
    data=json.loads(spec.read_text(encoding='utf-8'))
    for c in data['chapters']:
        if f"{c['chapter']:02d}"==number: return root/data['output_dir']/c['file']
    return None

def set_alt_text(shape,text):
    # Word reads the drawing's docPr/@descr as the image's alternative text.
    # Without it a screen reader announces every figure as an unnamed picture.
    shape._inline.docPr.set('descr',text)

def chapter_cover_alt(root,number):
    data=json.loads((root/'config/chapter-covers.json').read_text(encoding='utf-8'))
    for c in data['chapters']:
        if f"{c['chapter']:02d}"==number: return f"Chapter {c['chapter']} cover: {c['title']}. {c['focus']}"
    return f"Chapter {int(number)} cover"

def diagram_alt(root,key,caption):
    for d in json.loads((root/'config/diagrams.json').read_text(encoding='utf-8'))['diagrams']:
        if d['id']==key and d.get('description'): return d['description']
    return caption

def add_chapter_cover(doc,path,alt=None):
    # One US Letter cover page per chapter, set like the book cover: the page
    # holds only the image, followed by a page break into the chapter itself.
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.paragraph_format.page_break_before=True
    shape=p.add_run().add_picture(str(path),width=Inches(7.05))
    if alt: set_alt_text(shape,alt)
    # The chapter heading that follows carries page_break_before, so the cover page holds
    # only the image: a separate break paragraph here could spill onto an empty page.

def add_diagram(doc,root,key,caption,cfg):
    path=root/'assets/diagrams'/f'{key}.png'
    if not path.exists(): raise FileNotFoundError(f'Missing diagram {path}')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2); shape=p.add_run().add_picture(str(path),width=Inches(cfg['design']['diagram_width_inches']))
    set_alt_text(shape,diagram_alt(root,key,caption))
    c=doc.add_paragraph(style='Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_inline(c,caption,cfg)

def normalize_zip(path:Path):
    tmp=path.with_suffix('.normalized.docx')
    with ZipFile(path,'r') as zin, ZipFile(tmp,'w',ZIP_DEFLATED,compresslevel=9) as zout:
        for name in sorted(zin.namelist()):
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0)); info.compress_type=ZIP_DEFLATED; info.external_attr=0o600<<16; info.create_system=3; zout.writestr(info,zin.read(name))
    tmp.replace(path)

def build(root:Path,out:Path,template_only=False):
    global INDEX
    cfg=load(root); C=cfg['design']['colors']; doc=Document(); configure(doc,cfg)
    doc.core_properties.title=cfg['book']['title']; doc.core_properties.subject='Enterprise general intelligence: governed agentic systems architecture'; doc.core_properties.author=cfg['book']['author']; doc.core_properties.last_modified_by=cfg['book']['author']; doc.core_properties.keywords='agentic AI, enterprise architecture, developer platform, AI agents, LLM gateway, security, evaluation'; doc.core_properties.comments='Generated deterministically from book/manuscript.md; do not hand-edit the release DOCX.'; doc.core_properties.created=FIXED_DT; doc.core_properties.modified=FIXED_DT
    if template_only:
        add_paragraph(doc,'Book Style Template','Title',cfg) if False else None
        p=doc.add_paragraph(style='Title'); p.add_run('Book Style Template')
        for i,name in enumerate(['Heading 1','Heading 2','Heading 3','Heading 4'],1): doc.add_paragraph(f'{name} sample',style=name)
        add_paragraph(doc,'Normal paragraph with **bold**, *italic*, `inline code`, [R01], and a [link](https://example.com).',cfg)
        add_table(doc,[['Column A','Column B'],['Value','Value']],cfg); add_code(doc,'python','def example():\n    return True',cfg)
    else:
        cover=root/'assets/cover/cover-v1.0.0.png'
        if not cover.exists(): subprocess.run([sys.executable,str(root/'tools/generate_cover.py'),'--root',str(root)],check=True)
        md=(root/cfg['source']['canonical_markdown']).read_text(encoding='utf-8')
        after_part=False; title_seen=False; lead_keep_next=False; blocks=0
        parsed=list(parse(md)); list_num_id=None
        INDEX=book_index.Tracker(book_index.load_terms(root,md),book_index.excluded_units(root))
        in_index=False; index_field=False
        for block_index,b in enumerate(parsed):
            blocks+=1; kind=b[0]
            if kind=='cover':
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); shape=p.add_run().add_picture(str(cover),width=Inches(7.05)); set_alt_text(shape,f"Book cover: {cfg['book']['title']}, {cfg['book']['author']}"); continue
            if kind=='part': INDEX.part(); add_part(doc,b[1],b[2],b[3],cfg); after_part=True; lead_keep_next=False; continue
            if kind=='heading':
                level,text=b[1],b[2]
                INDEX.heading(level,text)
                if level==1: in_index=text==book_index.INDEX_HEADING
                if level==1 and not title_seen:
                    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.page_break_before=True; add_inline(p,text,cfg); title_seen=True; continue
                style=f'Heading {level}'
                cover=None
                if level==1:
                    num=re.match(r'^(\d{2}) — ',text); cover=chapter_cover_path(root,num.group(1)) if num else None
                    if cover is not None:
                        if not cover.exists(): raise FileNotFoundError(f'Missing chapter cover {cover}')
                        add_chapter_cover(doc,cover,chapter_cover_alt(root,num.group(1)))
                p=doc.add_paragraph(style=style); add_inline(p,text,cfg); p.paragraph_format.keep_with_next=True
                if level==1:
                    is_chapter=bool(re.match(r'^\d{2} — ',text) or text.startswith('Appendix '))
                    p.paragraph_format.page_break_before=True  # every chapter, appendix and front-matter H1 starts its own page
                    after_part=False
                lead_keep_next=level in (2,3)
                continue
            if kind=='diagram': add_diagram(doc,root,b[1],b[2],cfg); lead_keep_next=False; continue
            if kind=='code':
                next_block=parsed[block_index+1] if block_index+1<len(parsed) else None
                no_spacer=bool(next_block and next_block[0]=='para' and next_block[1].startswith('**How to know the repair worked.**'))
                add_code(doc,b[1],b[2],cfg,spacer=not no_spacer); lead_keep_next=False; continue
            if kind=='table':
                next_block=parsed[block_index+1] if block_index+1<len(parsed) else None
                no_spacer=bool(next_block and next_block[0]=='para' and next_block[1].startswith('**How to know the repair worked.**'))
                add_table(doc,b[1],cfg,spacer=not no_spacer); lead_keep_next=False; continue
            if kind=='bullet':
                p=add_paragraph(doc,b[1],cfg,'List Bullet'); p.paragraph_format.keep_together=True; mark_index(p,b[1])
                if lead_keep_next: p.paragraph_format.keep_with_next=True
                lead_keep_next=False; continue
            if kind=='number':
                p=add_paragraph(doc,b[1],cfg,'List Number'); p.paragraph_format.keep_together=True; mark_index(p,b[1])
                prev=parsed[block_index-1][0] if block_index else None
                if prev!='number' or list_num_id is None: list_num_id=new_numbered_list(doc)
                set_num(p,list_num_id)
                if lead_keep_next: p.paragraph_format.keep_with_next=True
                lead_keep_next=False; continue
            if kind=='quote':
                style='Book Evidence' if b[1].startswith('**Evidence ') else 'Book Callout'; p=add_paragraph(doc,b[1],cfg,style); mark_index(p,b[1]); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.right_indent=Inches(.12); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7); lead_keep_next=False; continue
            if kind=='para':
                txt=b[1]
                if txt.startswith('Book v1.0.0') or txt=='**'+cfg['book']['author']+'**': p=add_paragraph(doc,txt,cfg,'Subtitle',WD_ALIGN_PARAGRAPH.CENTER)
                elif txt.startswith('A structured enterprise reference') and title_seen: p=add_paragraph(doc,txt,cfg,'Normal',WD_ALIGN_PARAGRAPH.CENTER)
                else:
                    style='Book Reference' if txt.startswith('* **[R') else 'Normal'; p=add_paragraph(doc,txt,cfg,style); mark_index(p,txt)
                    if in_index and not index_field and re.fullmatch(r'\*\*[A-Z]\*\*',txt): open_index_field(p); index_field=True
                if lead_keep_next: p.paragraph_format.keep_with_next=True
                lead_keep_next=False
                continue
        if index_field: doc.paragraphs[-1]._p.append(field_run(None,'end'))
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out); normalize_zip(out)
    report={'output':str(out.relative_to(root)),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'bytes':out.stat().st_size,'template_only':template_only}
    if not template_only:
        from docx import Document as D
        check=D(out); report.update({'paragraphs':len(check.paragraphs),'tables':len(check.tables),'inline_images':len(check.inline_shapes),'chapter_headings':sum(1 for p in check.paragraphs if p.style.name=='Heading 1' and re.match(r'^\d{2} — ',p.text)),'appendix_headings':sum(1 for p in check.paragraphs if p.style.name=='Heading 1' and p.text.startswith('Appendix '))})
        (root/'qa/build-report.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2)); return report

if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]); ap.add_argument('--output',type=Path); ap.add_argument('--template-only',action='store_true'); a=ap.parse_args(); root=a.root.resolve(); default=root/('templates/book-style-template.docx' if a.template_only else 'dist/Enterprise-General-Intelligence-v1.0.0.docx'); build(root,(a.output or default).resolve(),a.template_only)
